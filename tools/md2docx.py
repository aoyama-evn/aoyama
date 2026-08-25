"""Markdown -> .docx cho bo tai lieu Mockup.

Cach dung:
    pip install python-docx
    python tools/md2docx.py docs/Mockup docs/Mockup/docx

Chay lai moi khi sua file .md de ban .docx khong bi lech.

Ho tro: heading, doan van, in dam/nghieng/code inline, link, bullet & numbered list,
checkbox list, bang GFM, blockquote, duong ke ngang, code fence (giu nguyen khoang
trang, tu dong chon co chu de vua kho giay).
"""
import os
import re
import sys
import unicodedata

_libs = os.path.join(os.path.dirname(os.path.abspath(__file__)), "libs")
if os.path.isdir(_libs):
    sys.path.insert(0, _libs)

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_ASCII = "Segoe UI"      # tieng Viet day du dau
BODY_EA = "Yu Gothic"        # tieng Nhat trong bang/wireframe
MONO = "MS Gothic"           # monospace that su, halfwidth/fullwidth chuan

NAVY = RGBColor(0x1A, 0x2E, 0x44)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
CODE_COLOR = RGBColor(0x1A, 0x1A, 0x1A)

USABLE_MM = 170.0  # A4 rong 210mm, le 2cm moi ben


def disp_width(s):
    """Be rong hien thi tinh theo o ky tu monospace (CJK = 2)."""
    w = 0
    for ch in s:
        if unicodedata.east_asian_width(ch) in ("W", "F"):
            w += 2
        else:
            w += 1
    return w


def set_run_font(run, ascii_font, ea_font, size_pt, bold=False, italic=False,
                 color=None):
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:cs"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)


def shade(element, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    element.append(el)


def para_shading(p, hex_fill):
    shade(p._element.get_or_add_pPr(), hex_fill)


def cell_shading(cell, hex_fill):
    shade(cell._element.get_or_add_tcPr(), hex_fill)


def left_bar(p, hex_color):
    ppr = p._element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hex_color)
    borders.append(left)
    ppr.append(borders)


INLINE = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))"
)


def add_inline(p, text, size=10.5, base_bold=False):
    """Tach **dam**, *nghieng*, `code`, [link](url) roi them vao doan van."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, BODY_ASCII, BODY_EA, size, bold=True)
        elif len(part) > 2 and part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            set_run_font(r, BODY_ASCII, BODY_EA, size, bold=base_bold, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            set_run_font(r, MONO, MONO, size - 0.5, bold=base_bold)
            shade(r._element.get_or_add_rPr(), "EFEFEF")
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+?)\]\(([^)]+?)\)", part)
            label, url = m.group(1), m.group(2)
            r = p.add_run(label)
            if url.startswith("http"):
                set_run_font(r, BODY_ASCII, BODY_EA, size, bold=base_bold,
                             color=RGBColor(0x15, 0x65, 0xC0))
                r.font.underline = True
            else:
                set_run_font(r, BODY_ASCII, BODY_EA, size, bold=base_bold,
                             italic=True, color=NAVY)
        else:
            r = p.add_run(part)
            set_run_font(r, BODY_ASCII, BODY_EA, size, bold=base_bold)


def add_code_block(doc, lines, lang):
    if lang == "mermaid":
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(2)
        r = cap.add_run("Sơ đồ (mã nguồn Mermaid — mở bản .md trên GitHub để xem "
                        "sơ đồ đồ họa)")
        set_run_font(r, BODY_ASCII, BODY_EA, 8.5, italic=True, color=GREY)

    width = max([disp_width(ln) for ln in lines] or [1])
    # be rong mot o monospace ~ 0.6 * co chu; quy ra mm de vua USABLE_MM
    size = 803.0 / max(width, 1)
    size = max(6.0, min(9.0, size))

    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.2)
        para_shading(p, "F7F7F7")
        r = p.add_run(ln if ln else " ")
        set_run_font(r, MONO, MONO, size, color=CODE_COLOR)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = True
    for i, cell_text in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, cell_text, size=9.5, base_bold=True)
        cell_shading(c, "1A2E44")
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i in range(len(header)):
            c = cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, row[i] if i < len(row) else "", size=9.5)
            if ri % 2 == 1:
                cell_shading(c, "F5F5F5")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_sep_row(line):
    return bool(re.match(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$", line.strip()))


def convert(md_path, out_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = BODY_ASCII
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_EA)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for attr in ("left_margin", "right_margin"):
        setattr(sec, attr, Cm(2.0))
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Hệ thống quản lý dịch vụ xe máy AOYAMA — Tài liệu Mockup v1.0")
    set_run_font(fr, BODY_ASCII, BODY_EA, 8, color=GREY)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # code fence
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block, lang)
            continue

        # bang
        if stripped.startswith("|") and i + 1 < n and is_sep_row(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # duong ke ngang
        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            p = doc.add_paragraph()
            ppr = p._element.get_or_add_pPr()
            b = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single")
            bot.set(qn("w:sz"), "6")
            bot.set(qn("w:space"), "1")
            bot.set(qn("w:color"), "D0D0D0")
            b.append(bot)
            ppr.append(b)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            if level == 1:
                pf.space_before = Pt(0)
                pf.space_after = Pt(14)
                size, color = 20, NAVY
            elif level == 2:
                pf.space_before = Pt(16)
                pf.space_after = Pt(6)
                size, color = 14.5, NAVY
                pf.page_break_before = False
            elif level == 3:
                pf.space_before = Pt(12)
                pf.space_after = Pt(4)
                size, color = 12, NAVY
            else:
                pf.space_before = Pt(10)
                pf.space_after = Pt(3)
                size, color = 11, NAVY
            pf.keep_with_next = True
            r = p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", text))
            set_run_font(r, BODY_ASCII, BODY_EA, size, bold=True, color=color)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.5)
            pf.space_before = Pt(6)
            pf.space_after = Pt(6)
            left_bar(p, "C8102E")
            para_shading(p, "FDF5F6")
            add_inline(p, " ".join(x.strip() for x in buf if x.strip()), size=10)
            for r in p.runs:
                r.italic = True
            continue

        # checkbox
        m = re.match(r"^(\s*)-\s+\[( |x|X)\]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.6 + 0.6 * indent)
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            box = p.add_run("☐  " if m.group(2) == " " else "☑  ")
            set_run_font(box, MONO, MONO, 11)
            add_inline(p, m.group(3), size=10.5)
            i += 1
            continue

        # bullet
        m = re.match(r"^(\s*)[-*•]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet")
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8 + 0.5 * indent)
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            add_inline(p, m.group(2), size=10.5)
            i += 1
            continue

        # numbered
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Number")
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8 + 0.5 * indent)
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            add_inline(p, m.group(3), size=10.5)
            i += 1
            continue

        # doan van
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.25
        add_inline(p, stripped, size=10.5)
        i += 1

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    src_dir, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)
    names = sorted(f for f in os.listdir(src_dir) if f.endswith(".md"))
    for name in names:
        base = os.path.splitext(name)[0]
        if base == "README":
            base = "00-Muc-luc"
        out = os.path.join(out_dir, base + ".docx")
        convert(os.path.join(src_dir, name), out)
        print("OK  %-34s -> %s" % (name, os.path.basename(out)))
